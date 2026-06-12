from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from django.contrib.auth.models import Group, Permission, User
from django.contrib.contenttypes.models import ContentType
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from docx import Document

from competencies.models import (
    Competence,
    CompetenceIndicator,
    CompetenceIndicatorImport,
    DisciplineCompetence,
)
from competencies.services import (
    IndicatorDocumentConverter,
    IndicatorImportError,
    IndicatorImportService,
)
from competencies.services.indicator_parser import IndicatorDocxParser
from core.models import AcademicDegree, AcademicTitle, CompetenceType, EducationLevel
from programs.models import EducationalProgram, ProgramProfile, TrainingDirection
from teachers.models import Department, Teacher


@pytest.fixture()
def indicator_schema():
    if connection.vendor != 'sqlite':
        pytest.skip('Тесты импорта индикаторов создают временную схему только в sqlite.')

    models = [
        ContentType,
        Permission,
        Group,
        User,
        EducationLevel,
        CompetenceType,
        AcademicDegree,
        AcademicTitle,
        Department,
        Teacher,
        TrainingDirection,
        ProgramProfile,
        EducationalProgram,
        Competence,
        CompetenceIndicatorImport,
        CompetenceIndicator,
        DisciplineCompetence,
    ]
    with connection.schema_editor() as schema_editor:
        for model in models:
            schema_editor.create_model(model)
    yield
    with connection.schema_editor() as schema_editor:
        for model in reversed(models):
            schema_editor.delete_model(model)


def _program_with_competence(code='ПК-1'):
    level = EducationLevel.objects.create(name='бакалавриат')
    direction = TrainingDirection.objects.create(
        education_level=level,
        code='09.03.02',
        name='Информационные системы и технологии',
    )
    profile = ProgramProfile.objects.create(
        training_direction=direction,
        code='09.03.02.01',
        name='Информационные системы и технологии в бизнесе',
    )
    department = Department.objects.create(
        number='22',
        short_name='ИС',
        full_name='Кафедра информационных систем',
    )
    program = EducationalProgram.objects.create(
        program_profile=profile,
        department=department,
        admission_year=2026,
    )
    competence_type = CompetenceType.objects.create(name=code.split('-', 1)[0])
    competence = Competence.objects.create(
        educational_program=program,
        competence_type=competence_type,
        code=code,
        name='Способен решать профессиональные задачи',
    )
    return program, competence


def _indicator_docx(*, competence_code='ПК-1', first_text='основы анализа данных') -> bytes:
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'Код и наименование профессиональной компетенции выпускника'
    table.cell(0, 1).text = 'Индикаторы достижения компетенции'
    table.cell(1, 0).text = f'{competence_code} Способен решать профессиональные задачи'
    table.cell(1, 1).text = (
        f'{competence_code}.1. Знает: {first_text}. '
        f'{competence_code}.2. Умеет: применять методы анализа. '
        f'{competence_code}.3. Владеет: навыками решения задач.'
    )
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(content, name='indicators.docx'):
    return SimpleUploadedFile(
        name,
        content,
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )


def test_parser_splits_multiple_indicators_from_one_table_cell():
    parsed = IndicatorDocxParser().parse_upload(_upload(_indicator_docx()))

    assert parsed.tables_found == 1
    assert [row.indicator_code for row in parsed.rows] == ['ПК-1.1', 'ПК-1.2', 'ПК-1.3']
    assert parsed.rows[0].competence_code == 'ПК-1'
    assert parsed.rows[0].source_table_number == 1
    assert parsed.rows[0].source_row_number == 2
    assert parsed.rows[0].text == 'Знает: основы анализа данных'


def test_doc_upload_uses_conversion_layer(monkeypatch):
    converter = IndicatorDocumentConverter()
    monkeypatch.setattr(
        converter,
        'convert_doc_bytes',
        lambda **kwargs: _indicator_docx(),
    )

    prepared = converter.prepare_upload(_upload(b'old-binary-word', name='indicators.doc'))

    assert prepared.source_filename == 'indicators.doc'
    assert prepared.original_content == b'old-binary-word'
    assert prepared.docx_content.startswith(b'PK')


def test_libreoffice_conversion_uses_generated_source_docx(monkeypatch):
    converter = IndicatorDocumentConverter()
    monkeypatch.setattr(
        'competencies.services.indicator_document_converter.tempfile.gettempdir',
        lambda: str(Path.cwd()),
    )
    monkeypatch.setattr(
        'competencies.services.indicator_document_converter.shutil.which',
        lambda command: '/usr/bin/soffice' if command == 'soffice' else None,
    )

    def fake_run(command, **kwargs):
        output_dir = Path(command[command.index('--outdir') + 1])
        (output_dir / 'source.docx').write_bytes(_indicator_docx())
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(
        'competencies.services.indicator_document_converter.subprocess.run',
        fake_run,
    )

    converted = converter.convert_doc_bytes(content=b'old-binary-word', source_filename='indicators.doc')

    assert converted.startswith(b'PK')


def test_repeated_import_skips_duplicates_and_updates_changed_text(indicator_schema):
    program, competence = _program_with_competence()
    service = IndicatorImportService()

    first = service.import_upload(_upload(_indicator_docx()), educational_program=program)
    repeated = service.import_upload(_upload(_indicator_docx()), educational_program=program)
    changed = service.import_upload(
        _upload(_indicator_docx(first_text='обновлённые основы анализа данных')),
        educational_program=program,
    )

    assert first.created_count == 3
    assert repeated.created_count == 0
    assert repeated.skipped_count == 3
    assert changed.updated_count == 1
    assert CompetenceIndicator.objects.filter(competence=competence).count() == 3
    assert CompetenceIndicator.objects.get(competence=competence, code='ПК-1.1').text == (
        'Знает: обновлённые основы анализа данных'
    )


def test_unknown_competence_blocks_rows_and_records_failed_batch(indicator_schema):
    program, _ = _program_with_competence()
    service = IndicatorImportService()

    with pytest.raises(IndicatorImportError, match='Импорт отменён') as captured:
        service.import_upload(
            _upload(_indicator_docx(competence_code='ПК-99')),
            educational_program=program,
        )

    batch = CompetenceIndicatorImport.objects.get(pk=captured.value.batch_id)
    assert batch.status == CompetenceIndicatorImport.Status.FAILED
    assert batch.error_count == 3
    assert 'ПК-99' in batch.error_summary
    assert not CompetenceIndicator.objects.exists()


def test_corrupted_word_file_records_failed_batch(indicator_schema):
    program, _ = _program_with_competence()
    service = IndicatorImportService()

    with pytest.raises(IndicatorImportError, match='повреждён') as captured:
        service.import_upload(
            _upload(b'not-a-docx'),
            educational_program=program,
        )

    batch = CompetenceIndicatorImport.objects.get(pk=captured.value.batch_id)
    assert batch.status == CompetenceIndicatorImport.Status.FAILED
    assert batch.error_count == 1
    assert 'повреждён' in batch.error_summary
